import base64
import binascii
import csv
import io
import json
from pathlib import Path
from typing import Any

from app.data.equipment import (
    NUMERIC_CATALOG_FIELDS,
    CatalogValidationError,
    merge_equipment_catalog,
    normalize_catalog_item,
)
from app.schemas import CatalogImportRequest, CatalogImportResult, Issue

REQUIRED_COLUMNS = {"type", "name", "role", "systems"}
NUMERIC_COLUMNS = NUMERIC_CATALOG_FIELDS


def import_catalog(request: CatalogImportRequest) -> CatalogImportResult:
    issues: list[Issue] = []
    try:
        raw = base64.b64decode(request.content_base64)
    except (binascii.Error, ValueError) as exc:
        return CatalogImportResult(
            imported=0,
            total=0,
            replaced=request.replace,
            issues=[Issue(level="error", code="CATALOG_IMPORT_DECODE_ERROR", message=str(exc))],
        )
    suffix = Path(request.filename).suffix.lower()

    try:
        if suffix == ".json":
            items = _parse_json(raw)
        elif suffix == ".csv":
            items = _parse_csv(raw)
        elif suffix in {".xlsx", ".xlsm"}:
            items = _parse_xlsx(raw)
        else:
            return CatalogImportResult(
                imported=0,
                total=0,
                replaced=request.replace,
                issues=[Issue(level="error", code="UNSUPPORTED_CATALOG_FILE", message="Only .json, .csv, .xlsx and .xlsm are supported.")],
            )
    except (UnicodeDecodeError, json.JSONDecodeError, ValueError) as exc:
        return CatalogImportResult(
            imported=0,
            total=0,
            replaced=request.replace,
            issues=[Issue(level="error", code="CATALOG_IMPORT_PARSE_ERROR", message=str(exc))],
        )

    valid_items = []
    for index, item in enumerate(items, start=1):
        if not isinstance(item, dict):
            issues.append(Issue(level="error", code="CATALOG_ROW_INVALID", message=f"Row {index} must be an object."))
            continue
        missing = REQUIRED_COLUMNS - set(item)
        if missing:
            issues.append(Issue(level="error", code="CATALOG_ROW_MISSING_COLUMNS", message=f"Row {index} misses columns: {', '.join(sorted(missing))}."))
            continue
        try:
            valid_items.append(_normalize_import_row(item))
        except ValueError as exc:
            issues.append(Issue(level="error", code="CATALOG_ROW_INVALID", message=f"Row {index}: {exc}"))

    if not any(issue.level == "error" for issue in issues):
        try:
            merge_equipment_catalog(valid_items, replace=request.replace)
        except CatalogValidationError as exc:
            issues.append(Issue(level="error", code="CATALOG_ROW_INVALID", message=str(exc)))

    return CatalogImportResult(imported=len(valid_items) if not any(issue.level == "error" for issue in issues) else 0, total=len(items), replaced=request.replace, issues=issues)


def _parse_json(raw: bytes) -> list[Any]:
    data = json.loads(raw.decode("utf-8-sig"))
    if isinstance(data, dict) and "items" in data:
        data = data["items"]
    if not isinstance(data, list):
        # 必须保持 ValueError：调用方在 `except (UnicodeDecodeError,
        # json.JSONDecodeError, ValueError)` 中统一转成结构化的
        # CATALOG_IMPORT_PARSE_ERROR（HTTP 200）。换成 TypeError 会越过该捕获，
        # 变成 500，并被 tests/test_catalog_management.py 的解析错误用例抓到。
        raise ValueError(  # noqa: TRY004
            "Catalog JSON must be a list or an object with an items list."
        )
    # Keep malformed rows intact so row validation can report their position;
    # coercing dict(42) here would escape as a TypeError/HTTP 500.
    return data


def _parse_csv(raw: bytes) -> list[dict[str, Any]]:
    text = raw.decode("utf-8-sig")
    reader = csv.DictReader(io.StringIO(text))
    return [dict(row) for row in reader]


def _parse_xlsx(raw: bytes) -> list[dict[str, Any]]:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise ValueError("XLSX import requires openpyxl.") from exc

    workbook = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
    sheet = workbook.active
    if sheet is None:
        # 没有活动工作表时原先会抛 AttributeError，绕过上层的 ValueError 捕获变成 500。
        raise ValueError("XLSX file has no active worksheet.")
    rows = list(sheet.iter_rows(values_only=True))
    if not rows:
        return []
    headers = [str(value).strip() if value is not None else "" for value in rows[0]]
    items = []
    for row in rows[1:]:
        if not any(value is not None and value != "" for value in row):
            continue
        item = {headers[index]: row[index] for index in range(min(len(headers), len(row))) if headers[index]}
        items.append(item)
    return items


def _normalize_import_row(item: dict[str, Any]) -> dict[str, Any]:
    for key in REQUIRED_COLUMNS:
        if item.get(key) in ("", None):
            raise ValueError(f"{key} is empty")
    return normalize_catalog_item(item)
